"""Freeze Wisconsin DHS statewide senior-care directories as a privacy-minimal snapshot."""

from __future__ import annotations

import argparse
import hashlib
import io
import json
from datetime import datetime, timezone
from pathlib import Path

import requests
from openpyxl import load_workbook

BASE = "https://www.dhs.wisconsin.gov/guide/"
SOURCES = {
    "adult-family-home": "afhdirexcel.xlsx",
    "community-based-residential-facility": "cbrfdirexcel.xlsx",
    "residential-care-apartment-complex": "rcacdirexcel.xlsx",
    "nursing-home": "nhdirexcel.xlsx",
    "hospice": "hospicedirexcel.xlsx",
    "home-health-agency": "hhadirexcel.xlsx",
}
OUT = Path(__file__).resolve().parents[1] / "apps/web/src/data/wisconsin-public-snapshot.json"


def cell(value: object) -> str:
    return str(value).strip() if value is not None else ""


def build() -> dict:
    sources = []
    rows = []
    retrieved = datetime.now(timezone.utc).isoformat()
    for source_class, filename in SOURCES.items():
        url = BASE + filename
        response = requests.get(url, timeout=30)
        response.raise_for_status()
        raw = response.content
        if not raw.startswith(b"PK"):
            raise ValueError(f"Not an XLSX file: {url}")
        sheet = load_workbook(io.BytesIO(raw), read_only=True, data_only=True).active
        values = sheet.values
        header = [cell(x) for x in next(values)]
        index = {name: i for i, name in enumerate(header) if name}
        class_rows = []
        for record in values:
            def get(name: str) -> str:
                return cell(record[index[name]]) if name in index else ""

            if source_class in {"adult-family-home", "community-based-residential-facility", "residential-care-apartment-complex"}:
                license_id, name = get("Facility ID"), get("Facility Name")
                ccn, city, capacity, kind = "", get("City"), get("Capacity"), get("Class")
                issued = get("Date Regular Issued") or get("Date Probationary Issued")
            elif source_class == "nursing-home":
                license_id, name = get("LIC#"), get("PROVIDER")
                ccn, city, capacity, kind, issued = get("CERTIFICATION"), get("CITY"), get("BEDCNT"), "", ""
            else:
                license_id, name = get("Lic #"), get("Facility")
                ccn, city, capacity, kind, issued = get("Provider #"), "", "", "", ""
            if not license_id.isdigit() or not name:
                continue
            if not (len(ccn) == 6 and ccn.isdigit()):
                ccn = ""
            class_rows.append({"class": source_class, "license": license_id, "name": name,
                               "city": city, "capacity": capacity, "subclass": kind,
                               "issued": issued, "ccn": ccn})
        sources.append({"class": source_class, "url": url, "sheet": sheet.title,
                        "httpLastModified": response.headers.get("Last-Modified"),
                        "sha256": hashlib.sha256(raw).hexdigest(), "rows": len(class_rows),
                        "distinctLicenses": len({r["license"] for r in class_rows}),
                        "rowsWithCcn": sum(bool(r["ccn"]) for r in class_rows)})
        rows.extend(class_rows)
    survey_url = BASE + "bal-monthly-additions.xlsx"
    survey_response = requests.get(survey_url, timeout=30)
    survey_response.raise_for_status()
    survey_raw = survey_response.content
    survey_values = load_workbook(io.BytesIO(survey_raw), read_only=True, data_only=True).active.values
    survey_header = [cell(x) for x in next(survey_values)]
    if survey_header != ["License ID", "Name", "Address", "City", "County", "Survey ID", "Letter Type"]:
        raise ValueError("Survey additions schema changed")
    residential_ids = {r["license"] for r in rows if r["class"] in
                       {"adult-family-home", "community-based-residential-facility", "residential-care-apartment-complex"}}
    surveys = [{"license": cell(v[0]), "name": cell(v[1]), "surveyId": cell(v[5]),
                "letterType": cell(v[6]), "exactRosterMatch": cell(v[0]) in residential_ids}
               for v in survey_values if cell(v[0]) and cell(v[5])]
    return {"regulator": "Wisconsin Department of Health Services, Division of Quality Assurance",
            "retrievedAt": retrieved, "generatedAt": datetime.now(timezone.utc).isoformat(), "sources": sources,
            "cmsSourceKeys": sum(bool(r["ccn"]) for r in rows), "cmsExactBridges": 0,
            "surveyAdditionsSource": {"url": survey_url, "sha256": hashlib.sha256(survey_raw).hexdigest(),
                                      "httpLastModified": survey_response.headers.get("Last-Modified"),
                                      "rows": len(surveys), "exactRosterMatches": sum(s["exactRosterMatch"] for s in surveys)},
            "surveyAdditions": surveys,
            "inspectionIndex": "PARTIAL_MONTHLY_ADDITIONS", "inspectionExactAttachments": sum(s["exactRosterMatch"] for s in surveys),
            "enforcementRows": "NOT_ACQUIRED", "enforcementExactAttachments": 0,
            "nameOnlyAdverseJoins": 0, "graphWrites": 0,
            "newCanonicalFacilities": 0, "claimEligibilityChanges": 0,
            "rows": rows}


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--check", action="store_true")
    args = parser.parse_args()
    snapshot = build()
    existing = json.loads(OUT.read_text(encoding="utf-8")) if OUT.exists() else None
    if args.check:
        if not existing:
            raise SystemExit("Snapshot missing")
        for old, new in zip(existing["sources"], snapshot["sources"], strict=True):
            if old["sha256"] != new["sha256"]:
                raise SystemExit(f"Source changed: {old['class']}")
        if existing["surveyAdditionsSource"]["sha256"] != snapshot["surveyAdditionsSource"]["sha256"]:
            raise SystemExit("Survey additions source changed")
        print("Wisconsin DHS source hashes match frozen snapshot")
        return
    OUT.write_text(json.dumps(snapshot, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    for source in snapshot["sources"]:
        print(source["class"], source["rows"], source["distinctLicenses"], source["rowsWithCcn"], source["httpLastModified"])


if __name__ == "__main__":
    main()

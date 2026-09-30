"""Build a frozen, location-grain TX HHSC review packet. No database access or writes."""

from __future__ import annotations

import csv
import hashlib
import json
from collections import Counter
from datetime import date, datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
BASE = ROOT / "data/enrichment/th-enrich-b2/idr2"
RAW = BASE / "raw"
RELEASE = BASE / "release"
BATCH = "TH-ENRICH-TX-SENIOR-2026-09-30-IDR2"
AS_OF = "2026-09-29"
SOURCES = (
    ("ICFIID.xlsx", "TX_ICF_IID", 708),
    ("DAHS.xlsx", "TX_DAHS", 389),
    ("dahs_issonly.xlsx", "TX_DAHS_ISS_ONLY", 743),
)


def cell(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip() if value is not None else ""


def key(provider_class: str, facility_id: str) -> str:
    if provider_class not in {source[1] for source in SOURCES} or not facility_id or not facility_id.isdigit():
        raise ValueError("Invalid Texas HHSC class or Facility ID")
    return f"TX|HHSC|{provider_class}|{facility_id}"


def read_sheet(path: Path):
    sheet = load_workbook(path, read_only=True, data_only=True).active
    values = iter(sheet.values)
    title = cell(next(values)[0])
    headers = [cell(item) for item in next(values)]
    rows = [(index, dict(zip(headers, map(cell, row)))) for index, row in enumerate(values, 3)
            if any(cell(item) for item in row)]
    dimensions = (sheet.max_row, sheet.max_column)
    sheet.parent.close()
    return title, rows, *dimensions


def write_csv(path: Path, rows: list[dict]):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def build():
    providers, licenses, manifest = [], [], []
    prior_path = ROOT / "data/enrichment/th-enrich-b2/raw/dahs_issonly.xlsx"
    _, prior_rows, _, _ = read_sheet(prior_path)
    prior_112184 = [row for _, row in prior_rows if row["Facility ID"] == "112184"]
    if len(prior_112184) != 1:
        raise ValueError("Expected one historical 112184 row in prior 744-row source")
    for filename, provider_class, expected in SOURCES:
        path = RAW / filename
        title, rows, worksheet_rows, worksheet_columns = read_sheet(path)
        if len(rows) != expected or f"as of 09/29/2026" not in title:
            raise ValueError(f"{filename}: live source count/date drift; stop")
        if provider_class == "TX_DAHS_ISS_ONLY":
            old_ids = {row["Facility ID"] for _, row in prior_rows}
            current_ids = {row["Facility ID"] for _, row in rows}
            if old_ids - current_ids != {"112184"} or current_ids - old_ids:
                raise ValueError("In-home source changed beyond the one Evidence-identified ID")
        if any(row["Facility ID"] == "112184" for _, row in rows):
            raise ValueError(f"{filename}: rejected Facility ID 112184 reappeared; review drift")
        body = path.read_bytes()
        source_hash = hashlib.sha256(body).hexdigest()
        source_url = f"https://apps.hhs.texas.gov/providers/directories/{filename}"
        retrieved_at = datetime.fromtimestamp(path.stat().st_mtime, timezone.utc).isoformat()
        manifest.append({"file": filename, "official_url": source_url,
                         "sha256": source_hash, "bytes": len(body), "worksheet_rows": worksheet_rows,
                         "worksheet_columns": worksheet_columns, "source_rows": len(rows), "class": provider_class,
                         "as_of": AS_OF,
                         "retrieved_at": retrieved_at,
                         "source_title": title})
        for source_row, raw in rows:
            facility_id = raw["Facility ID"]
            namespaced_key = key(provider_class, facility_id)
            provider = {"batch_id": BATCH, "namespaced_key": namespaced_key,
                        "provider_class": provider_class, "facility_id": facility_id,
                        "source_row": source_row, "official_name": raw["Facility Name"],
                        "program_type": raw["Program Type"],
                        "facility_licensed_raw": raw["Facility  Licensed"],
                        "facility_certified_raw": raw["Facility Certified"],
                        "county": raw["County"], "physical_address": raw["Physical Address"],
                        "city": raw["Physical Address CITY"], "state": raw["Physical Address State"],
                        "zip": raw["Physical Address Zipcode"], "source_sha256": source_hash,
                        "source_url": source_url, "retrieved_at": retrieved_at,
                        "source_as_of": AS_OF, "organization_id": "", "public_eligible": "NO"}
            provider["record_sha256"] = hashlib.sha256(json.dumps(provider, sort_keys=True,
                                                separators=(",", ":")).encode()).hexdigest()
            providers.append(provider)
            if raw["License No"]:
                licenses.append({"batch_id": BATCH, "namespaced_key": namespaced_key,
                                 "credential_type": "HHSC_LICENSE_NUMBER", "credential_number": raw["License No"],
                                 "facility_licensed_raw": raw["Facility  Licensed"],
                                 "license_effective_date": raw["License Effective Date"],
                                 "license_expiration_date": raw["License Expiration Date"],
                                 "source_sha256": source_hash})
    keys = [row["namespaced_key"] for row in providers]
    if len(keys) != 1840 or len(set(keys)) != 1840 or len(licenses) != 1786:
        raise ValueError("Texas location/license invariants failed")
    if sum(row["facility_licensed_raw"] == "NO" for row in licenses) != 2:
        raise ValueError("Unlicensed license-number observation drift")
    if any(row["facility_id"] == "112184" for row in providers):
        raise ValueError("Prior-only Facility ID survived")
    write_csv(RELEASE / "locations.csv", providers)
    write_csv(RELEASE / "license_observations.csv", licenses)
    receipt = {"batch_id": BATCH, "source_as_of": AS_OF,
               "source_rows": dict(Counter(row["provider_class"] for row in providers)),
               "total_locations": len(providers), "license_observations": len(licenses),
               "location_keys_unique": len(set(keys)), "cross_class_key_collisions": len(keys) - len(set(keys)),
               "status_by_class": {
                   provider_class: {
                       "facility_licensed": dict(Counter(row["facility_licensed_raw"] for row in providers
                                                         if row["provider_class"] == provider_class)),
                       "facility_certified": dict(Counter(row["facility_certified_raw"] for row in providers
                                                          if row["provider_class"] == provider_class)),
                       "license_observations": sum(row["namespaced_key"].split("|")[2] == provider_class
                                                   for row in licenses),
                   } for _, provider_class, _ in SOURCES},
               "source_manifest": manifest,
               "prior_112184": {"present_prior": True, "present_now": False,
                                "prior_facility_name": prior_112184[0]["Facility Name"],
                                "prior_license_number": prior_112184[0]["License No"],
                                "reason_for_removal": "Not stated in the official workbooks"},
               "files": {p.name: hashlib.sha256(p.read_bytes()).hexdigest()
                         for p in (RELEASE / "locations.csv", RELEASE / "license_observations.csv")},
               "production_mutations": False}
    (BASE / "receipt.json").write_text(json.dumps(receipt, indent=2) + "\n", encoding="utf-8")
    return receipt


if __name__ == "__main__":
    print(json.dumps(build(), indent=2))

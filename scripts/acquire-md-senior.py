"""Freeze public OHCQ facility directory workbooks without contact information."""
import hashlib
import io
import json
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from urllib.request import Request, urlopen

from openpyxl import load_workbook

BASE = "https://health.maryland.gov/ohcq/docs/Provider-Listings/Excel/"
FILES = {
    "Assisted Living Programs": "Assisted-Living-Excel.xlsx",
    "Long Term Care Facilities": "Long-Term-Care-Facilities-EXCEL.xlsx",
    "Home Health Agencies": "Home-Health-Agencies-EXCEL.xlsx",
    "Hospices": "Hospice-EXCEL.xlsx",
    "Adult Medical Day Care Centers": "Adult-Medical-Day-Care-Centers-EXCEL.xlsx",
}
OUT = Path(__file__).resolve().parents[1] / "apps/web/src/data/maryland-public-snapshot.json"


def value(v):
    return str(v).strip() if v is not None else ""


def main():
    sources, rows = [], []
    for kind, filename in FILES.items():
        url = BASE + filename
        with urlopen(Request(url, headers={"User-Agent": "SeniorTrustHub research contact info@seniortrusthub.com"}), timeout=40) as response:
            data = response.read()
            modified = response.headers.get("Last-Modified")
        sheet = load_workbook(io.BytesIO(data), read_only=True, data_only=True).active
        cells = sheet.iter_rows(values_only=True)
        headers = [value(x).lower() for x in next(cells)]
        source_rows = []
        for raw in cells:
            record = dict(zip(headers, raw))
            name = value(record.get("licensee"))
            if not name:
                continue
            row = {
                "class": kind,
                "license": value(record.get("license number")),
                "name": name,
                "city": value(record.get("city")),
                "state": value(record.get("state")),
                "capacity": record.get("capacity") if isinstance(record.get("capacity"), int) else None,
                "levelOfCare": value(record.get("level of care")) or None,
                "certificationType": value(record.get("certiftype")) or None,
            }
            rows.append(row)
            source_rows.append(row)
        licenses = [r["license"] for r in source_rows if r["license"]]
        sources.append({
            "class": kind, "url": url, "sheet": sheet.title,
            "httpLastModified": modified, "sha256": hashlib.sha256(data).hexdigest(),
            "rows": len(source_rows), "distinctLicenses": len(set(licenses)),
            "missingLicenseRows": len(source_rows) - len(licenses),
            "duplicateLicenseRows": sum(n - 1 for n in Counter(licenses).values() if n > 1),
        })
    output = {"regulator": "Maryland Department of Health, Office of Health Care Quality",
              "retrievedAt": datetime.now(timezone.utc).isoformat(),
              "sources": sources, "rows": rows,
              "cmsExactBridges": 0, "inspectionAttachments": 0, "enforcementExactAttachments": 0,
              "newCanonicalFacilities": 0, "graphWrites": 0, "claimEligibilityChanges": 0}
    OUT.write_text(json.dumps(output, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(json.dumps(sources, indent=2))


if __name__ == "__main__":
    main()

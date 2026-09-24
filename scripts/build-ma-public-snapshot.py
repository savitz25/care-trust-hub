#!/usr/bin/env python3
"""MA-SEN-001 — Massachusetts senior-care state layer over the accepted CMS spine.

Stage 1 (--parse; needs gitignored raw files in data/massachusetts/ma-sen-001/raw/):
  DPH "List of health care facilities licensed or certified by the Division" (XLSX) -> senior classes
  AGE "Certified Assisted Living Residences" list (XLSX) -> sanitized residence rows (no staff names,
  emails, or federal tax IDs)
  Writes data/massachusetts/ma-sen-001/*.json (committed).
Stage 2 (default): derived JSON + accepted CMS national partition -> public snapshot, TS module,
  and the server-side facility list used by /massachusetts. --check compares the committed files.

Does not download CMS. Does not use MassGIS as a roster. No name-only state-to-CMS joins.
"""

from __future__ import annotations

import hashlib
import json
import re
import sys
from collections import Counter
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data/massachusetts/ma-sen-001/raw"
STAGE = ROOT / "data/massachusetts/ma-sen-001"
ART = ROOT / "artifacts/ma-sen-001-public-snapshot.json"
TS = ROOT / "packages/domain/src/ma-public-snapshot.ts"
WEB_LIST = ROOT / "apps/web/src/data/massachusetts-facility-lists.json"
NATIONAL = json.loads(
    (ROOT / "apps/web/src/data/senior-national-intelligence.json").read_text(
        encoding="utf-8"
    )
)
GENERATION_KEYS = frozenset({"generatedAt", "fingerprint"})
DPH_RETRIEVED_AT = "2026-09-24T16:37:44Z"
ALR_RETRIEVED_AT = "2026-09-24T16:37:52Z"
CENSUS_RETRIEVED_AT = "2026-09-24T16:37:59Z"
SURVEY_TOOL_OBSERVED_AT = "2026-09-24T16:38:59Z"
GENERATED_AT = "2026-09-24T17:00:00Z"
SNAPSHOT_AS_OF = "2026-09-24"
DPH_PAGE = "https://www.mass.gov/info-details/find-information-about-licensed-or-certified-health-care-facilities"
DPH_FILE = "https://www.mass.gov/doc/list-of-health-care-facilities-licensed-or-certified-by-the-division/download"
ALR_PAGE = "https://www.mass.gov/assisted-living-residences"
ALR_FILE = "https://mass.gov/doc/list-of-certified-assisted-living-residences-as-of-april-2025/download"
ALR_DIRECTORY = "https://www.mass.gov/assisted-living-residences/locations"
CENSUS_FILE = (
    "https://www.mass.gov/doc/assisted-living-residences-census-report-2026/download"
)
SURVEY_TOOL = "https://eohhs.ehs.state.ma.us/nursehome/Default.aspx"
SENIOR_CLASSES = {
    "Nursing Home": "nursingHomes",
    "Rest Home": "restHomes",
    "Certified Home Health Agency": "homeHealth",
    "Hospice": "hospice",
    "Hospice Inpatient Satellite": "hospiceInpatientSatellites",
    "Adult Day Health": "adultDayHealth",
}
CITIES = ["BOSTON", "WORCESTER", "SPRINGFIELD"]


def dumps(obj: object) -> str:
    return json.dumps(obj, sort_keys=True, separators=(",", ":"), ensure_ascii=False)


def strip_generation(obj: object) -> object:
    if isinstance(obj, dict):
        return {
            k: strip_generation(v) for k, v in obj.items() if k not in GENERATION_KEYS
        }
    if isinstance(obj, list):
        return [strip_generation(v) for v in obj]
    return obj


def sha(obj: dict) -> str:
    return hashlib.sha256(dumps(strip_generation(obj)).encode("utf-8")).hexdigest()


def file_sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def clean(v: object) -> str:
    return re.sub(r"\s+", " ", str(v if v is not None else "")).strip()


def as_int(v: object) -> int | None:
    s = clean(v)
    return int(float(s)) if re.fullmatch(r"\d+(\.0+)?", s) else None


def parse() -> None:
    import openpyxl

    ws = openpyxl.load_workbook(
        RAW / "dph-licensed-certified-facilities.xlsx", read_only=True
    )
    rows = list(ws.worksheets[0].iter_rows(values_only=True))
    source_date = rows[6][0]
    if not isinstance(source_date, datetime):
        raise SystemExit(f"DPH workbook date cell moved: {rows[6]}")
    header = [clean(a) + " " + clean(b) for a, b in zip(rows[8], rows[9])]
    expected = [
        "Type of Facility",
        "DPH Facility ID Number",
        "Name of Facility",
        "Street",
        "City/Town",
        "Zip Code",
        "Telephone",
        "Bed Count (if applicable)",
        "Adult Day Health Capacity (if applicable)",
    ]
    if [h.strip() for h in header] != expected:
        raise SystemExit(f"DPH header drifted: {header}")
    body = [r for r in rows[10:] if any(c is not None for c in r)]
    all_types = Counter(clean(r[0]) for r in body)
    facilities = []
    for r in body:
        kind = clean(r[0])
        if kind not in SENIOR_CLASSES:
            continue
        facilities.append(
            {
                "class": SENIOR_CLASSES[kind],
                "typeOfFacility": kind,
                "dphFacilityId": clean(r[1]),
                "name": clean(r[2]),
                "street": clean(r[3]),
                "city": clean(r[4]),
                "zip": clean(r[5]),
                "telephone": clean(r[6]),
                "beds": as_int(r[7]),
                "adultDayHealthCapacity": as_int(r[8]),
            }
        )
    ids = Counter(f["dphFacilityId"] for f in facilities)
    if any(v > 1 for v in ids.values()):
        raise SystemExit("duplicate DPH facility id across senior classes")
    facilities.sort(key=lambda f: (f["class"], f["name"], f["dphFacilityId"]))
    (STAGE / "dph-senior-facilities.json").write_text(
        json.dumps(
            {
                "source": DPH_FILE,
                "sourcePage": DPH_PAGE,
                "sourceTitle": "Massachusetts Licensed or Certified Health Care Facility/Agency Listing",
                "sourceAsOf": source_date.date().isoformat(),
                "retrievedAt": DPH_RETRIEVED_AT,
                "sha256": file_sha(RAW / "dph-licensed-certified-facilities.xlsx"),
                "workbookRowsAllTypes": len(body),
                "workbookTypes": dict(sorted(all_types.items())),
                "columns": expected,
                "facilities": facilities,
            },
            indent=1,
            ensure_ascii=False,
        )
        + "\n",
        encoding="utf-8",
    )

    ws = openpyxl.load_workbook(
        RAW / "age-certified-alr-list.xlsx", read_only=True, data_only=True
    )
    rows = list(ws.worksheets[0].iter_rows(values_only=True))
    title = clean(rows[0][0])
    m = re.search(r"as of\s+([A-Za-z]+ \d{1,2}, \d{4})", title)
    if not m:
        raise SystemExit(f"ALR title has no as-of date: {title}")
    alr_as_of = (
        datetime.strptime(m.group(1) + " +0000", "%B %d, %Y %z").date().isoformat()
    )
    h = {clean(k): i for i, k in enumerate(rows[1]) if k}
    residences = []
    for r in rows[2:]:
        if not r[0]:
            continue
        status = r[h["Status"]]
        residences.append(
            {
                "name": clean(r[h["ALR Name"]]),
                "statusAsPublished": clean(status)
                if not isinstance(status, datetime)
                else f"(date in status column: {status.date().isoformat()})",
                "street": clean(r[h["Address"]]),
                "city": clean(r[h["City"]]),
                "county": clean(r[h["County"]]),
                "telephone": clean(r[h["Telephone"]]),
                "initialCertification": clean(r[h["Initial certification"]]),
                "totalUnits": as_int(r[h["Total # of Units"]]),
                "traditionalUnits": as_int(r[h["# Traditional Units"]]),
                "specialCareUnits": as_int(r[h["# SCR Units"]]),
                "maxOccupancy": as_int(r[h["Total Max Occupancy"]]),
            }
        )
    residences.sort(key=lambda a: (a["name"], a["street"]))
    (STAGE / "age-alr-residences.json").write_text(
        json.dumps(
            {
                "source": ALR_FILE,
                "sourcePage": ALR_PAGE,
                "sourceTitle": title,
                "sourceAsOf": alr_as_of,
                "retrievedAt": ALR_RETRIEVED_AT,
                "sha256": file_sha(RAW / "age-certified-alr-list.xlsx"),
                "fieldsWithheld": [
                    "Executive Director",
                    "ED email",
                    "RCD NAME",
                    "RCD Email",
                    "RCD History",
                    "ED Change History",
                    "additional e-mails",
                    "Federal ID #",
                    "Staff updating",
                ],
                "residences": residences,
            },
            indent=1,
            ensure_ascii=False,
        )
        + "\n",
        encoding="utf-8",
    )


def clock(sources: dict, key: str) -> dict:
    row = sources[key]
    return {
        "datasetKey": row["datasetKey"],
        "officialUrl": row["officialUrl"],
        "sourceModifiedAt": row["sourceModifiedAt"],
        "retrievedAt": row["retrievedAt"],
        "sourcePeriod": row["sourcePeriod"],
    }


def dph_class(facilities: list[dict], key: str, dph: dict, **extra: object) -> dict:
    rows = [f for f in facilities if f["class"] == key]
    beds = [f["beds"] for f in rows if f["beds"] is not None]
    return {
        "rows": len(rows),
        "distinctDphFacilityIds": len({f["dphFacilityId"] for f in rows}),
        "bedsSum": sum(beds) if beds else None,
        "rowsWithBeds": len(beds),
        "source": "DPH Licensed or Certified Health Care Facility/Agency Listing",
        "sourceAsOf": dph["sourceAsOf"],
        "retrievedAt": dph["retrievedAt"],
        "capability": "KNOWN",
        "statusField": "NONE_IN_SOURCE",
        "ccnInSource": False,
        **extra,
    }


def build() -> tuple[dict, dict]:
    dph = json.loads((STAGE / "dph-senior-facilities.json").read_text(encoding="utf-8"))
    alr = json.loads((STAGE / "age-alr-residences.json").read_text(encoding="utf-8"))
    facilities = dph["facilities"]
    residences = alr["residences"]
    cms = next(row for row in NATIONAL["geography"] if row["state"] == "MA")
    sources = {row["datasetKey"]: row for row in NATIONAL["sources"]}
    if cms["nursingHomes"] != 341 or cms["homeHealth"] != 287 or cms["hospice"] != 77:
        raise SystemExit(f"Massachusetts CMS partition drifted: {cms}")

    def units(key: str) -> int:
        return sum(a[key] for a in residences if a[key] is not None)

    city_counts = {
        city.title(): {
            "nursingHomes": sum(
                1
                for f in facilities
                if f["class"] == "nursingHomes" and f["city"] == city
            ),
            "restHomes": sum(
                1 for f in facilities if f["class"] == "restHomes" and f["city"] == city
            ),
            # AGE writes Boston neighborhoods as "Boston-Dorchester", "Boston-Brighton", ...
            "assistedLivingResidences": sum(
                1
                for a in residences
                if a["city"].upper() == city or a["city"].upper().startswith(city + "-")
            ),
            "homeHealth": sum(
                1
                for f in facilities
                if f["class"] == "homeHealth" and f["city"] == city
            ),
            "hospice": sum(
                1 for f in facilities if f["class"] == "hospice" and f["city"] == city
            ),
        }
        for city in CITIES
    }
    snapshot = {
        "version": "senior-ma-state-intel-v1",
        "ticket": "MA-SEN-001",
        "publicationPath": "/massachusetts",
        "asOf": SNAPSHOT_AS_OF,
        "snapshotAsOf": SNAPSHOT_AS_OF,
        "retrievedAt": DPH_RETRIEVED_AT,
        "generatedAt": GENERATED_AT,
        "localWorkNeededNow": "NO",
        "claimEligibilityBroadened": False,
        "no_trust_score": True,
        "no_aggregate_rating": True,
        "no_rankings": True,
        "no_boston_page": True,
        "clocks": {
            "generatedAt": GENERATED_AT,
            "snapshotAsOf": SNAPSHOT_AS_OF,
            "dph_workbook_sourceAsOf": dph["sourceAsOf"],
            "dph_workbook_retrievedAt": dph["retrievedAt"],
            "alr_list_sourceAsOf": alr["sourceAsOf"],
            "alr_list_retrievedAt": alr["retrievedAt"],
            "alr_census_report_period": "Calendar year 2025; data as of 2025-12-31",
            "alr_census_report_retrievedAt": CENSUS_RETRIEVED_AT,
            "survey_tool_processed_through": "2026-09-04",
            "survey_tool_observedAt": SURVEY_TOOL_OBSERVED_AT,
            "cms_nursing_home_sourceAsOf": sources["nursing-home-provider-information"][
                "sourceModifiedAt"
            ],
            "cms_home_health_sourceAsOf": sources["home-health-care-agencies"][
                "sourceModifiedAt"
            ],
            "cms_hospice_sourceAsOf": sources["hospice-general-information"][
                "sourceModifiedAt"
            ],
            "license_effective_date": None,
            "license_expiration_date": None,
            "retrievedAt_is_not_sourceAsOf": True,
            "cms_date_is_not_massachusetts_license_date": True,
        },
        "regulatorMap": {
            "dph": "Massachusetts Department of Public Health, Division of Health Care Facility Licensure and Certification",
            "age": "Massachusetts Executive Office of Aging & Independence (AGE)",
            "dphFacilityList": DPH_PAGE,
            "alrProgram": ALR_PAGE,
            "alrDirectory": ALR_DIRECTORY,
            "alrReports": "https://www.mass.gov/lists/annual-assisted-living-residence-alr-data-reports",
            "surveyTool": SURVEY_TOOL,
            "surveyToolGuide": "https://www.mass.gov/guides/nursing-home-survey-performance-tool",
            "nursingHomeConsumerInfo": "https://www.mass.gov/nursing-home-consumer-information",
            "closures": "https://www.mass.gov/info-details/information-about-nursing-home-closures",
            "cmsCareCompare": "https://www.medicare.gov/care-compare/",
        },
        "cmsOverlay": {
            "nursingHomes": cms["nursingHomes"],
            "homeHealth": cms["homeHealth"],
            "hospice": cms["hospice"],
            "source": "senior-national-intelligence.json geography MA (CMS class directories)",
            "asOf": "2026-08-27",
            "nationalFingerprint": NATIONAL["sourceFingerprint"],
            "addedToNationalTotals": False,
            "censusConfirmation": {
                "artifact": "artifacts/senior-metric-census-r2-03.json",
                "nursingHomes": 341,
                "homeHealth": 287,
                "hospice": 77,
            },
            "clocks": {
                "nursingHomes": clock(sources, "nursing-home-provider-information"),
                "homeHealth": clock(sources, "home-health-care-agencies"),
                "hospice": clock(sources, "hospice-general-information"),
                "inspections": clock(sources, "nursing-home-inspection-dates"),
                "ownership": clock(sources, "skilled-nursing-facility-all-owners"),
                "penalties": clock(sources, "nursing-home-penalties"),
            },
        },
        "dphWorkbook": {
            "source": DPH_FILE,
            "sourceAsOf": dph["sourceAsOf"],
            "retrievedAt": dph["retrievedAt"],
            "sha256": dph["sha256"],
            "rowsAllTypes": dph["workbookRowsAllTypes"],
            "typesInWorkbook": len(dph["workbookTypes"]),
            "seniorTypesTaken": list(SENIOR_CLASSES),
            "otherTypesExcluded": sorted(
                t for t in dph["workbookTypes"] if t not in SENIOR_CLASSES
            ),
            "columns": dph["columns"],
            "ccnInSource": False,
            "statusField": "NONE_IN_SOURCE",
        },
        "nursingHomes": dph_class(
            facilities,
            "nursingHomes",
            dph,
            distinctFromRestHome=True,
            distinctFromCmsNursingHome=True,
        ),
        "restHomes": dph_class(
            facilities,
            "restHomes",
            dph,
            definition="Supportive residential setting for people who need 24-hour supervision but not routine nursing or medical care.",
            distinctFromNursingHome=True,
            distinctFromAssistedLiving=True,
        ),
        "homeHealth": dph_class(
            facilities,
            "homeHealth",
            dph,
            stateTypeLabel="Certified Home Health Agency",
            distinctFromCmsHomeHealth=True,
        ),
        "hospice": dph_class(facilities, "hospice", dph, distinctFromCmsHospice=True),
        "hospiceInpatientSatellites": dph_class(
            facilities,
            "hospiceInpatientSatellites",
            dph,
            subUnitOfHospiceProgram=True,
            notAddedToHospiceRows=True,
        ),
        "adultDayHealth": dph_class(
            facilities,
            "adultDayHealth",
            dph,
            capacitySum=sum(
                f["adultDayHealthCapacity"] or 0
                for f in facilities
                if f["class"] == "adultDayHealth"
            )
            or None,
            notResidential=True,
        ),
        "assistedLiving": {
            "rows": len(residences),
            "statusCounts": dict(Counter(a["statusAsPublished"] for a in residences)),
            "totalUnits": units("totalUnits"),
            "traditionalUnits": units("traditionalUnits"),
            "specialCareUnits": units("specialCareUnits"),
            "maxOccupancy": units("maxOccupancy"),
            "residencesWithSpecialCareUnits": sum(
                1 for a in residences if (a["specialCareUnits"] or 0) > 0
            ),
            "unitGrain": "units reported on the AGE list; not residents and not beds",
            "regulator": "Executive Office of Aging & Independence (AGE)",
            "credential": "ALR certification (not a DPH license, not CMS certification)",
            "source": ALR_FILE,
            "sourceTitle": alr["sourceTitle"],
            "sourceAsOf": alr["sourceAsOf"],
            "retrievedAt": alr["retrievedAt"],
            "sha256": alr["sha256"],
            "fieldsWithheld": alr["fieldsWithheld"],
            "liveDirectoryObservation": {
                "results": 273,
                "observedAt": ALR_RETRIEVED_AT,
                "url": ALR_DIRECTORY,
                "rowsAcquired": None,
                "note": "The live Mass.gov location listing showed 273 results; the downloadable list (as of the date above) has the rows used here.",
            },
            "capability": "KNOWN",
            "distinctFromRestHome": True,
            "distinctFromNursingHome": True,
        },
        "alrCensus2026": {
            "source": CENSUS_FILE,
            "retrievedAt": CENSUS_RETRIEVED_AT,
            "reportPeriod": "Calendar year 2025 (data as of 2025-12-31)",
            "statedCertifiedAlrsJanuary2026": 272,
            "respondingAlrs": 268,
            "usedAsIdentitySource": False,
            "capability": "KNOWN",
        },
        "surveyTool": {
            "name": "DPH Nursing Home Survey Performance Tool",
            "url": SURVEY_TOOL,
            "processedThrough": "2026-09-04",
            "listedFacilities": 343,
            "observedAt": SURVEY_TOOL_OBSERVED_AT,
            "method": "Tool's own methodology: 132 items reviewed across the last 3 standard surveys (44 per survey). Not a TrustHub score.",
            "facilityIdentifiersInListing": False,
            "indexedFacilityResults": None,
            "capability": "KNOWN",
            "bulkCapability": "NOT_ACQUIRED",
            "reason": "Results are reached through paged ASP.NET postbacks and per-facility pages with truncated names and no IDs. Not backfilled.",
            "notTrustHubScore": True,
        },
        "complaints": {
            "intake": "KNOWN",
            "providerLevelRows": None,
            "capability": "REQUEST_ONLY",
            "complaintIsNotDeficiency": True,
        },
        "closures": {
            "page": "https://www.mass.gov/info-details/information-about-nursing-home-closures",
            "rows": None,
            "capability": "NOT_ACQUIRED",
            "reason": "Closure notices are individual PDF/Word filings; not parsed.",
        },
        "massgis": {
            "capability": "UNSUPPORTED",
            "reason": "MassGIS Long-Term Care Residences (March 2024 release, sources from late 2023) is not a current roster.",
        },
        "crosswalk": {
            "exactStateToCmsBridges": 0,
            "attempted": True,
            "method": "exact CCN in state source",
            "ccnInDphWorkbook": False,
            "nameOnly": "UNSAFE",
            "addressResolver": "NOT_USED (no accepted address resolver contract for state-to-CMS joins)",
            "dphNursingHomeRowsUnresolved": sum(
                1 for f in facilities if f["class"] == "nursingHomes"
            ),
            "cmsNursingHomeRowsUnresolved": cms["nursingHomes"],
            "capability": "UNKNOWN",
            "reason": "The DPH workbook has no CCN, so no exact bridge exists. Zero bridges is not zero overlap: most DPH nursing homes are likely CMS-certified, but that is not established here.",
        },
        "cityContext": {
            "grain": "facility rows with that city in the source address; not a local page",
            "cities": city_counts,
        },
        "expansionLedger": {
            "NET_NEW_STATE_RESEARCH_IDENTITIES": len(facilities) + len(residences),
            "NET_NEW_CANONICAL_FACILITIES": 0,
            "NET_NEW_PUBLIC_PROFILES": 0,
            "EXISTING_FACILITIES_ENRICHED": 0,
            "GRAPH_WRITES": 0,
            "CLAIM_ELIGIBILITY_BROADENED": False,
        },
        "capabilities": [
            {"id": "cms-nursing-home-ma", "state": "KNOWN"},
            {"id": "cms-home-health-ma", "state": "KNOWN"},
            {"id": "cms-hospice-ma", "state": "KNOWN"},
            {"id": "dph-facility-workbook", "state": "KNOWN"},
            {"id": "dph-nursing-home-rows", "state": "KNOWN"},
            {"id": "dph-rest-home-rows", "state": "KNOWN"},
            {"id": "dph-home-health-rows", "state": "KNOWN"},
            {"id": "dph-hospice-rows", "state": "KNOWN"},
            {"id": "dph-adult-day-health-rows", "state": "KNOWN"},
            {"id": "age-alr-list", "state": "KNOWN"},
            {"id": "age-alr-census-2026-context", "state": "KNOWN"},
            {"id": "dph-survey-performance-tool", "state": "KNOWN"},
            {"id": "dph-survey-results-bulk", "state": "NOT_ACQUIRED"},
            {"id": "dph-license-status-or-expiration", "state": "UNKNOWN"},
            {"id": "dph-provider-complaints", "state": "REQUEST_ONLY"},
            {"id": "nursing-home-closure-list", "state": "NOT_ACQUIRED"},
            {"id": "exact-state-to-cms-join", "state": "UNKNOWN"},
            {"id": "name-only-state-to-cms-join", "state": "UNSUPPORTED"},
            {"id": "massgis-as-current-roster", "state": "UNSUPPORTED"},
            {"id": "combined-massachusetts-senior-facilities", "state": "UNSUPPORTED"},
        ],
        "fingerprint": "",
    }
    snapshot["fingerprint"] = sha(snapshot)
    lists = {
        "fingerprint": snapshot["fingerprint"],
        "dphSourceAsOf": dph["sourceAsOf"],
        "alrSourceAsOf": alr["sourceAsOf"],
        "restHomes": [f for f in facilities if f["class"] == "restHomes"],
        "nursingHomes": [f for f in facilities if f["class"] == "nursingHomes"],
        "assistedLivingResidences": residences,
    }
    return snapshot, lists


def ts_module(body: dict) -> str:
    return (
        "/** Generated by scripts/build-ma-public-snapshot.py. Do not edit by hand. */\n"
        "export const MA_PUBLIC_SNAPSHOT = "
        + json.dumps(body, indent=2, ensure_ascii=False)
        + " as const;\nexport type MaPublicSnapshot = typeof MA_PUBLIC_SNAPSHOT;\n"
    )


def main() -> None:
    if "--parse" in sys.argv:
        parse()
    body, lists = build()
    art = json.dumps(body, indent=2, ensure_ascii=False) + "\n"
    web = json.dumps(lists, indent=1, ensure_ascii=False) + "\n"
    if "--check" in sys.argv:
        for path, want in ((ART, art), (WEB_LIST, web)):
            if path.read_text(encoding="utf-8").replace("\r\n", "\n") != want:
                raise SystemExit(f"{path.relative_to(ROOT)} drifted from the builder")
        committed_ts = TS.read_text(encoding="utf-8").replace("\r\n", "\n")
        if f'"fingerprint": "{body["fingerprint"]}"' not in committed_ts:
            raise SystemExit("ma-public-snapshot.ts fingerprint drifted")
        print("MA-SEN-001 snapshot check OK", body["fingerprint"])
        return
    ART.write_text(art, encoding="utf-8")
    WEB_LIST.write_text(web, encoding="utf-8")
    TS.write_text(ts_module(body), encoding="utf-8")
    print(
        json.dumps(
            {
                "fingerprint": body["fingerprint"],
                "nh": body["nursingHomes"]["rows"],
                "rest": body["restHomes"]["rows"],
                "hha": body["homeHealth"]["rows"],
                "hospice": body["hospice"]["rows"],
                "adh": body["adultDayHealth"]["rows"],
                "alr": body["assistedLiving"]["rows"],
                "units": body["assistedLiving"]["totalUnits"],
                "cities": body["cityContext"]["cities"],
            },
            indent=1,
        )
    )


if __name__ == "__main__":
    main()
